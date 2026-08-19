"""Synthetic-only workbook fixtures for G1 profiling tests."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill


TABLE_IDS = ("PRBD01N001", "PREF01N001", "PREF02N001", "PRFD01N001")
PHYSICAL_TOKENS = {
    "PRBD01N001": "PRBD01N001",
    "PREF01N001": "PREF01N001",
    "PREF02N001": "PREF02N001",
    "PRFD01N001": "PRFDN001",
}


@dataclass(frozen=True, slots=True)
class SyntheticWorkbookSet:
    """Paths and writers for a complete, entirely synthetic source set."""

    raw_dir: Path
    schema_dir: Path
    raw_paths: dict[str, Path]
    schema_paths: dict[str, Path]

    @property
    def all_paths(self) -> tuple[Path, ...]:
        return tuple(self.raw_paths.values()) + tuple(self.schema_paths.values())

    def write_data(
        self,
        table_id: str,
        *,
        headers: list[Any],
        rows: list[list[Any]],
        preamble: list[list[Any]] | None = None,
        trailing_blank_rows: int = 0,
        merged_ranges: tuple[str, ...] = (),
    ) -> Path:
        """Replace one synthetic datarows workbook with the requested shape."""

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "datarows"
        for row in preamble or []:
            sheet.append(row)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)

        # Styled empty cells survive XLSX serialization and model physical
        # trailing rows without introducing synthetic values.
        blank_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
        for _ in range(trailing_blank_rows):
            row_index = sheet.max_row + 1
            for column_index in range(1, max(len(headers), 1) + 1):
                sheet.cell(row=row_index, column=column_index).fill = blank_fill

        for merged_range in merged_ranges:
            sheet.merge_cells(merged_range)

        path = self.raw_paths[table_id]
        workbook.save(path)
        workbook.close()
        return path

    def write_schema(
        self,
        table_id: str,
        *,
        fields: list[str],
        field_header: str = "컬럼명",
        preamble: list[list[Any]] | None = None,
        duplicate_field_rows: list[str] | None = None,
        merged_ranges: tuple[str, ...] = (),
        examples: dict[str, Any] | None = None,
    ) -> Path:
        """Replace one synthetic schema workbook without financial semantics."""

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1_Schema"
        for row in preamble or []:
            sheet.append(row)
        sheet.append([field_header, "컬럼타입", "컬럼한글명", "컬럼값예시"])
        for field in [*fields, *(duplicate_field_rows or [])]:
            sheet.append(
                [
                    field,
                    "TEXT",
                    "synthetic test field",
                    (examples or {}).get(field),
                ]
            )
        for merged_range in merged_ranges:
            sheet.merge_cells(merged_range)

        path = self.schema_paths[table_id]
        workbook.save(path)
        workbook.close()
        return path


@pytest.fixture
def synthetic_workbooks(tmp_path: Path) -> SyntheticWorkbookSet:
    """Create exactly four raw and four schema workbooks under ``tmp_path``."""

    raw_dir = tmp_path / "raw"
    schema_dir = tmp_path / "schema"
    raw_dir.mkdir()
    schema_dir.mkdir()

    raw_paths = {
        table_id: raw_dir
        / f"{PHYSICAL_TOKENS[table_id]}_synthetic_20260711_datarows.xlsx"
        for table_id in TABLE_IDS
    }
    schema_paths = {
        table_id: schema_dir / f"{table_id}_synthetic_schema.xlsx"
        for table_id in TABLE_IDS
    }
    workbooks = SyntheticWorkbookSet(
        raw_dir=raw_dir,
        schema_dir=schema_dir,
        raw_paths=raw_paths,
        schema_paths=schema_paths,
    )

    for table_id in TABLE_IDS:
        workbooks.write_data(
            table_id,
            headers=["field_a", "field_b"],
            rows=[["alpha", "one"], ["beta", "two"]],
        )
        workbooks.write_schema(table_id, fields=["field_a", "field_b"])

    return workbooks
