from __future__ import annotations

import json
import re
from dataclasses import replace
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from pipeline.normalize.models import ParseStatus
from pipeline.normalize.type_rules import build_typed_rule
from pipeline.profiling.discovery import SourceWorkbook, sha256_file
from pipeline.structured.loader import (
    StagingLoadError,
    expected_catalog_columns,
    iter_typed_rows,
    load_all,
    load_table,
    render_insert_sql,
)
from pipeline.structured.schema import build_staging_schema, build_table_spec


def _source(path: Path, table_id: str) -> SourceWorkbook:
    stat = path.stat()
    return SourceWorkbook(
        source_kind="datarows",
        logical_table_id=table_id,
        path=path,
        physical_filename=path.name,
        extension=".xlsx",
        size_bytes=stat.st_size,
        modified_time_ns=stat.st_mtime_ns,
        modified_at_utc=datetime.fromtimestamp(
            stat.st_mtime, timezone.utc
        ).isoformat(),
        sha256=sha256_file(path),
    )


def _write_workbook(path: Path, headers: list[str], rows: list[list[Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "datarows"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def _profile(table_id: str, column_names: list[str], row_count: int) -> dict[str, Any]:
    return {
        "logical_table_id": table_id,
        "selected_data_sheet": "datarows",
        "header_row": 1,
        "last_data_row": row_count + 1,
        "row_count": row_count,
        "columns": [
            {
                "column_name": name,
                "occurrence": 1,
                "qualified_name": f"{name}#1",
            }
            for name in column_names
        ],
    }


def _rules(table_id: str):
    return [
        build_typed_rule(
            logical_table_id=table_id,
            source_column="id",
            declared_type="text",
        ),
        build_typed_rule(
            logical_table_id=table_id,
            source_column="amount",
            declared_type="numeric",
        ),
        build_typed_rule(
            logical_table_id=table_id,
            source_column="label",
            declared_type="text",
        ),
    ]


class FakeCursor:
    def __init__(self, connection: "FakeConnection") -> None:
        self.connection = connection

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, traceback):
        return None

    @staticmethod
    def _table(sql: str) -> str | None:
        match = re.search(r'"structured_staging"\."([a-z0-9_]+)"', sql)
        return match.group(1) if match else None

    def execute(self, query: str, params=None):
        self.connection.statements.append((query, params))
        if "FROM information_schema.columns" in query:
            self.connection.last_catalog_query = "columns"
            self.connection.last_selected_table = params[1]
            return
        if "FROM information_schema.table_constraints" in query:
            if "constraint_type = 'PRIMARY KEY'" in query:
                self.connection.last_catalog_query = "primary_key"
            elif "constraint_type = 'UNIQUE'" in query:
                self.connection.last_catalog_query = "unique"
            elif "information_schema.check_constraints" in query:
                self.connection.last_catalog_query = "check"
            else:
                raise AssertionError("unexpected constraint catalog query")
            self.connection.last_selected_table = params[1]
            return
        table = self._table(query)
        if query.startswith("DELETE FROM") and table:
            self.connection.rows[table] = []
        if query.startswith("SELECT count"):
            self.connection.last_catalog_query = None
            self.connection.last_selected_table = table

    def executemany(self, query: str, params_seq):
        params = list(params_seq)
        self.connection.statements.append((query, params))
        table = self._table(query)
        assert table is not None
        self.connection.rows.setdefault(table, []).extend(params)

    def fetchone(self):
        table = self.connection.last_selected_table
        rows = self.connection.rows.get(table or "", [])
        row_count = len(rows)
        distinct = len({row[0] for row in rows})
        if table == self.connection.fail_validation_table:
            distinct = max(0, distinct - 1)
        return row_count, distinct

    def fetchall(self):
        table = self.connection.last_selected_table or ""
        if self.connection.last_catalog_query == "columns":
            return self.connection.catalog_columns.get(table, ())
        if self.connection.last_catalog_query == "primary_key":
            return [(name,) for name in self.connection.primary_keys.get(table, ())]
        if self.connection.last_catalog_query == "unique":
            return [
                (constraint_name, column_name, position)
                for constraint_name, columns in self.connection.unique_constraints.get(
                    table, {}
                ).items()
                for position, column_name in enumerate(columns, start=1)
            ]
        if self.connection.last_catalog_query == "check":
            return [
                (clause,) for clause in self.connection.check_clauses.get(table, ())
            ]
        raise AssertionError("unexpected fetchall")


class FakeConnection:
    def __init__(self, *, fail_validation_table: str | None = None) -> None:
        self.rows: dict[str, list[tuple[Any, ...]]] = {}
        self.statements: list[tuple[str, Any]] = []
        self.last_selected_table: str | None = None
        self.last_catalog_query: str | None = None
        self.fail_validation_table = fail_validation_table
        self.catalog_columns: dict[str, tuple[tuple[str, str, str, int], ...]] = {}
        self.primary_keys: dict[str, tuple[str, ...]] = {}
        self.unique_constraints: dict[str, dict[str, tuple[str, ...]]] = {}
        self.check_clauses: dict[str, tuple[str, ...]] = {}
        self.commits = 0
        self.rollbacks = 0

    def register_table(self, table) -> None:
        self.catalog_columns[table.db_table_name] = expected_catalog_columns(table)
        self.primary_keys[table.db_table_name] = ("source_row_id",)
        self.unique_constraints[table.db_table_name] = {
            "lineage_unique": (
                "source_file_sha256",
                "source_sheet",
                "source_excel_row",
            )
        }
        self.check_clauses[table.db_table_name] = (
            f'(("logical_table_id") = (\'{table.logical_table_id}\'::text))',
            '(("source_excel_row") > (0))',
        )

    def cursor(self):
        return FakeCursor(self)

    def commit(self):
        self.commits += 1

    def rollback(self):
        self.rollbacks += 1


def test_iter_typed_rows_preserves_raw_parse_status_and_lineage(tmp_path: Path) -> None:
    path = tmp_path / "TEST01_datarows.xlsx"
    injected = "x'); DROP TABLE anything; --"
    _write_workbook(
        path,
        ["id", "amount", "label"],
        [
            ["A", "12.50", " keep "],
            ["B", "", injected],
            ["C", "not-a-number", None],
        ],
    )
    source = _source(path, "TEST01")
    rules = _rules("TEST01")
    table = build_table_spec("TEST01", rules)

    rows = list(
        iter_typed_rows(
            source=source,
            table=table,
            rules=rules,
            data_profile_table=_profile("TEST01", ["id", "amount", "label"], 3),
        )
    )

    assert len(rows) == len({row.lineage.source_row_id for row in rows}) == 3
    assert rows[0].typed_values == ("A", Decimal("12.50"), " keep ")
    assert rows[0].raw_record["label#1"] == {
        "type": "STRING",
        "value": " keep ",
    }
    assert rows[1].raw_record["amount#1"] == {"type": "STRING", "value": ""}
    assert rows[1].parse_statuses["amount#1"] == ParseStatus.EMPTY.value
    assert rows[1].raw_record["label#1"]["value"] == injected
    assert rows[2].typed_values[1] is None
    assert rows[2].parse_statuses["amount#1"] == ParseStatus.PARSE_ERROR.value
    assert path.stat().st_size == source.size_bytes
    assert sha256_file(path) == source.sha256


def test_load_table_batches_parameters_and_commits_after_validation(tmp_path: Path) -> None:
    path = tmp_path / "TEST01_datarows.xlsx"
    injected = "x'); DELETE FROM structured_staging; --"
    _write_workbook(
        path,
        ["id", "amount", "label"],
        [["A", "1.25", injected], ["B", "2", "two"], ["C", "3", "three"]],
    )
    source = _source(path, "TEST01")
    rules = _rules("TEST01")
    table = build_table_spec("TEST01", rules)
    connection = FakeConnection()
    connection.register_table(table)

    stats = load_table(
        connection,
        source=source,
        table=table,
        rules=rules,
        data_profile_table=_profile("TEST01", ["id", "amount", "label"], 3),
        batch_size=2,
    )

    assert stats.validated is True
    assert stats.batch_count == 2
    assert stats.inserted_row_count == stats.database_row_count == 3
    assert connection.commits == 1
    assert connection.rollbacks == 0
    inserts = [(sql, params) for sql, params in connection.statements if sql.startswith("INSERT")]
    assert len(inserts) == 2
    assert all("%s" in sql for sql, _ in inserts)
    assert all(injected not in sql for sql, _ in inserts)
    assert any(injected in params[0][5] for _, params in inserts)
    assert "::jsonb" in render_insert_sql(table)


def test_load_table_rolls_back_on_source_row_id_validation_failure(tmp_path: Path) -> None:
    path = tmp_path / "TEST01_datarows.xlsx"
    _write_workbook(path, ["id", "amount", "label"], [["A", "1", "one"]])
    source = _source(path, "TEST01")
    rules = _rules("TEST01")
    table = build_table_spec("TEST01", rules)
    connection = FakeConnection(fail_validation_table=table.db_table_name)
    connection.register_table(table)

    with pytest.raises(StagingLoadError, match="database validation failed"):
        load_table(
            connection,
            source=source,
            table=table,
            rules=rules,
            data_profile_table=_profile("TEST01", ["id", "amount", "label"], 1),
        )
    assert connection.commits == 0
    assert connection.rollbacks == 1


def test_load_all_is_one_transaction_for_all_four_product_tables(tmp_path: Path) -> None:
    table_ids = ("PRBD01N001", "PREF01N001", "PREF02N001", "PRFD01N001")
    all_rules = []
    sources = []
    profiles = []
    for index, table_id in enumerate(table_ids, start=1):
        path = tmp_path / f"{table_id}_datarows.xlsx"
        identifier = "" if table_id == "PREF01N001" else f"ID-{index}"
        _write_workbook(path, ["id", "amount", "label"], [[identifier, str(index), " x "]])
        sources.append(_source(path, table_id))
        rules = _rules(table_id)
        all_rules.extend(rules)
        profiles.append(_profile(table_id, ["id", "amount", "label"], 1))
    schema = build_staging_schema(
        all_rules,
        expected_column_counts={table_id: 3 for table_id in table_ids},
    )
    connection = FakeConnection()
    for table in schema.tables:
        connection.register_table(table)

    stats = load_all(
        connection,
        sources=sources,
        schema=schema,
        rules=all_rules,
        data_profile={"tables": profiles},
        batch_size=1,
    )

    assert stats.validated is True
    assert stats.source_row_count == stats.inserted_row_count == 4
    assert len(stats.tables) == 4
    assert connection.commits == 1
    assert connection.rollbacks == 0
    pref_row = connection.rows["stg_pref01n001"][0]
    raw_record = json.loads(pref_row[5])
    statuses = json.loads(pref_row[6])
    assert raw_record["id#1"] == {"type": "STRING", "value": ""}
    assert statuses["id#1"] == "EMPTY"


def test_load_all_rolls_back_every_table_when_late_validation_fails(tmp_path: Path) -> None:
    table_ids = ("PRBD01N001", "PREF01N001", "PREF02N001", "PRFD01N001")
    all_rules = []
    sources = []
    profiles = []
    for table_id in table_ids:
        path = tmp_path / f"{table_id}_datarows.xlsx"
        _write_workbook(path, ["id", "amount", "label"], [["ID", "1", "one"]])
        sources.append(_source(path, table_id))
        rules = _rules(table_id)
        all_rules.extend(rules)
        profiles.append(_profile(table_id, ["id", "amount", "label"], 1))
    schema = build_staging_schema(
        all_rules,
        expected_column_counts={table_id: 3 for table_id in table_ids},
    )
    connection = FakeConnection(fail_validation_table="stg_prfd01n001")
    for table in schema.tables:
        connection.register_table(table)

    with pytest.raises(StagingLoadError, match="database validation failed"):
        load_all(
            connection,
            sources=sources,
            schema=schema,
            rules=all_rules,
            data_profile={"tables": profiles},
        )
    assert connection.commits == 0
    assert connection.rollbacks == 1


def test_changed_source_is_rejected_before_database_write(tmp_path: Path) -> None:
    path = tmp_path / "TEST01_datarows.xlsx"
    _write_workbook(path, ["id", "amount", "label"], [["A", "1", "one"]])
    source = _source(path, "TEST01")
    rules = _rules("TEST01")
    table = build_table_spec("TEST01", rules)
    stale = replace(source, sha256="0" * 64)
    connection = FakeConnection()
    connection.register_table(table)

    with pytest.raises(StagingLoadError, match="changed before"):
        load_table(
            connection,
            source=stale,
            table=table,
            rules=rules,
            data_profile_table=_profile("TEST01", ["id", "amount", "label"], 1),
        )
    assert connection.commits == 0
    assert connection.rollbacks == 1


def test_incompatible_existing_catalog_is_rejected_before_delete(tmp_path: Path) -> None:
    path = tmp_path / "TEST01_datarows.xlsx"
    _write_workbook(path, ["id", "amount", "label"], [["A", "1", "one"]])
    source = _source(path, "TEST01")
    rules = _rules("TEST01")
    table = build_table_spec("TEST01", rules)
    connection = FakeConnection()
    connection.register_table(table)
    connection.catalog_columns[table.db_table_name] = connection.catalog_columns[
        table.db_table_name
    ][:-1]

    with pytest.raises(StagingLoadError, match="table shape is incompatible"):
        load_table(
            connection,
            source=source,
            table=table,
            rules=rules,
            data_profile_table=_profile("TEST01", ["id", "amount", "label"], 1),
        )
    assert connection.commits == 0
    assert connection.rollbacks == 1
    assert not any(sql.startswith("DELETE FROM") for sql, _ in connection.statements)


def test_missing_lineage_constraint_is_rejected_before_delete(tmp_path: Path) -> None:
    path = tmp_path / "TEST01_datarows.xlsx"
    _write_workbook(path, ["id", "amount", "label"], [["A", "1", "one"]])
    source = _source(path, "TEST01")
    rules = _rules("TEST01")
    table = build_table_spec("TEST01", rules)
    connection = FakeConnection()
    connection.register_table(table)
    connection.unique_constraints[table.db_table_name] = {}

    with pytest.raises(StagingLoadError, match="unique lineage constraint"):
        load_table(
            connection,
            source=source,
            table=table,
            rules=rules,
            data_profile_table=_profile("TEST01", ["id", "amount", "label"], 1),
        )
    assert connection.commits == 0
    assert connection.rollbacks == 1
    assert not any(sql.startswith("DELETE FROM") for sql, _ in connection.statements)
