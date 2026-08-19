"""Read-only XLSX to PostgreSQL typed-staging loader.

The loader is synchronous and accepts a psycopg-style connection object.  It
does not import a driver, open a network connection, or read ``DATABASE_URL``;
connection construction remains an application boundary.  This keeps the
contract testable with fakes and makes transaction ownership explicit.
"""

from __future__ import annotations

import json
import re
from collections import Counter
from collections.abc import Iterable, Iterator, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Protocol

from pipeline.normalize.lineage import SourceLineage, build_source_lineage
from pipeline.normalize.models import SourceScalar, TypedRule
from pipeline.normalize.normalizer import normalize_value
from pipeline.profiling.discovery import (
    SourceWorkbook,
    sha256_file,
    verify_sources_unchanged,
)
from pipeline.profiling.statistics import is_blank_value

from .schema import (
    StagingSchemaSpec,
    StagingTableSpec,
    quote_identifier,
    render_create_schema_sql,
    render_create_table_sql,
)


class CursorLike(Protocol):
    def execute(self, query: str, params: Sequence[Any] | None = None) -> Any: ...

    def executemany(self, query: str, params_seq: Sequence[Sequence[Any]]) -> Any: ...

    def fetchone(self) -> Any: ...

    def fetchall(self) -> Sequence[Any]: ...

    def __enter__(self) -> "CursorLike": ...

    def __exit__(self, exc_type: Any, exc: Any, traceback: Any) -> None: ...


class ConnectionLike(Protocol):
    def cursor(self) -> CursorLike: ...

    def commit(self) -> None: ...

    def rollback(self) -> None: ...


class StagingLoadError(RuntimeError):
    """Raised when source, rule, load, or post-load validation is unsafe."""


_CATALOG_COLUMNS_SQL = """
SELECT column_name, data_type, is_nullable, ordinal_position
FROM information_schema.columns
WHERE table_schema = %s AND table_name = %s
ORDER BY ordinal_position
""".strip()

_CATALOG_PRIMARY_KEY_SQL = """
SELECT kcu.column_name
FROM information_schema.table_constraints AS tc
JOIN information_schema.key_column_usage AS kcu
  ON tc.constraint_catalog = kcu.constraint_catalog
 AND tc.constraint_schema = kcu.constraint_schema
 AND tc.constraint_name = kcu.constraint_name
WHERE tc.table_schema = %s
  AND tc.table_name = %s
  AND tc.constraint_type = 'PRIMARY KEY'
ORDER BY kcu.ordinal_position
""".strip()

_CATALOG_UNIQUE_SQL = """
SELECT tc.constraint_name, kcu.column_name, kcu.ordinal_position
FROM information_schema.table_constraints AS tc
JOIN information_schema.key_column_usage AS kcu
  ON tc.constraint_catalog = kcu.constraint_catalog
 AND tc.constraint_schema = kcu.constraint_schema
 AND tc.constraint_name = kcu.constraint_name
WHERE tc.table_schema = %s
  AND tc.table_name = %s
  AND tc.constraint_type = 'UNIQUE'
ORDER BY tc.constraint_name, kcu.ordinal_position
""".strip()

_CATALOG_CHECKS_SQL = """
SELECT cc.check_clause
FROM information_schema.table_constraints AS tc
JOIN information_schema.check_constraints AS cc
  ON tc.constraint_catalog = cc.constraint_catalog
 AND tc.constraint_schema = cc.constraint_schema
 AND tc.constraint_name = cc.constraint_name
WHERE tc.table_schema = %s
  AND tc.table_name = %s
  AND tc.constraint_type = 'CHECK'
ORDER BY tc.constraint_name
""".strip()


@dataclass(frozen=True, slots=True)
class StagedRow:
    """One source-preserving database row before driver adaptation."""

    lineage: SourceLineage
    raw_record: Mapping[str, Mapping[str, Any]]
    parse_statuses: Mapping[str, str]
    typed_values: tuple[SourceScalar, ...]

    def parameters(self, table: StagingTableSpec) -> tuple[Any, ...]:
        if len(self.typed_values) != len(table.columns):
            raise StagingLoadError("typed value count does not match staging table")
        # Pass deterministic JSON text and cast placeholders to jsonb in SQL.
        # No driver-specific Jsonb adapter is needed.
        raw_json = json.dumps(
            self.raw_record,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            allow_nan=False,
        )
        statuses_json = json.dumps(
            self.parse_statuses,
            ensure_ascii=True,
            sort_keys=True,
            separators=(",", ":"),
        )
        return (
            self.lineage.source_row_id,
            self.lineage.source_table,
            self.lineage.source_excel_row,
            self.lineage.source_file_sha256,
            self.lineage.source_sheet,
            raw_json,
            statuses_json,
            *self.typed_values,
        )


@dataclass(frozen=True, slots=True)
class TableLoadStats:
    logical_table_id: str
    source_row_count: int
    inserted_row_count: int
    database_row_count: int
    distinct_source_row_id_count: int
    batch_count: int
    parse_status_counts: Mapping[str, int]
    source_unchanged: bool

    @property
    def validated(self) -> bool:
        return (
            self.source_row_count
            == self.inserted_row_count
            == self.database_row_count
            == self.distinct_source_row_id_count
            and self.source_unchanged
        )

    def to_dict(self) -> dict[str, Any]:
        return {
            "logical_table_id": self.logical_table_id,
            "source_row_count": self.source_row_count,
            "inserted_row_count": self.inserted_row_count,
            "database_row_count": self.database_row_count,
            "distinct_source_row_id_count": self.distinct_source_row_id_count,
            "batch_count": self.batch_count,
            "parse_status_counts": dict(self.parse_status_counts),
            "source_unchanged": self.source_unchanged,
            "validated": self.validated,
        }


@dataclass(frozen=True, slots=True)
class LoadStats:
    """Committed all-table load summary."""

    tables: tuple[TableLoadStats, ...]
    committed: bool

    @property
    def source_row_count(self) -> int:
        return sum(table.source_row_count for table in self.tables)

    @property
    def inserted_row_count(self) -> int:
        return sum(table.inserted_row_count for table in self.tables)

    @property
    def validated(self) -> bool:
        return self.committed and all(table.validated for table in self.tables)

    def to_dict(self) -> dict[str, Any]:
        return {
            "table_count": len(self.tables),
            "source_row_count": self.source_row_count,
            "inserted_row_count": self.inserted_row_count,
            "committed": self.committed,
            "validated": self.validated,
            "tables": [table.to_dict() for table in self.tables],
        }


def _raw_envelope(value: SourceScalar) -> dict[str, Any]:
    """Encode a source scalar without trimming, imputing, or guessing units."""

    if value is None:
        return {"type": "NULL", "value": None}
    if isinstance(value, bool):
        return {"type": "BOOLEAN", "value": "true" if value else "false"}
    if isinstance(value, str):
        return {"type": "STRING", "value": value}
    if isinstance(value, int):
        return {"type": "INTEGER", "value": str(value)}
    if isinstance(value, float):
        return {"type": "FLOAT", "value": value.hex()}
    if isinstance(value, Decimal):
        return {"type": "DECIMAL", "value": str(value)}
    if isinstance(value, datetime):
        return {"type": "DATETIME", "value": value.isoformat()}
    if isinstance(value, date):
        return {"type": "DATE", "value": value.isoformat()}
    if isinstance(value, time):
        return {"type": "TIME", "value": value.isoformat()}
    raise StagingLoadError(f"unsupported source scalar type: {type(value)!r}")


def _qualified_headers(values: tuple[Any, ...]) -> list[str]:
    end = len(values)
    while end and values[end - 1] is None:
        end -= 1
    counts: Counter[str] = Counter()
    qualified: list[str] = []
    for value in values[:end]:
        name = "<blank>" if value is None or str(value).strip() == "" else str(value).strip()
        counts[name] += 1
        qualified.append(f"{name}#{counts[name]}")
    return qualified


def _load_workbook(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - installation guidance
        raise RuntimeError(
            "structured staging requires: python -m pip install -e \".[data]\""
        ) from error
    return load_workbook(path, read_only=True, data_only=False, keep_links=False)


def _validate_rule_alignment(
    table: StagingTableSpec,
    rules: Sequence[TypedRule],
) -> None:
    expected = [column.qualified_name for column in table.columns]
    observed = [str(rule.qualified_name) for rule in rules]
    if expected != observed:
        raise StagingLoadError(
            f"typed-rule order differs from staging schema for {table.logical_table_id}"
        )
    if any(rule.logical_table_id != table.logical_table_id for rule in rules):
        raise StagingLoadError("typed rule belongs to a different logical table")


def iter_typed_rows(
    *,
    source: SourceWorkbook,
    table: StagingTableSpec,
    rules: Sequence[TypedRule],
    data_profile_table: Mapping[str, Any],
) -> Iterator[StagedRow]:
    """Yield deterministic typed rows without retaining the workbook in memory."""

    if source.source_kind != "datarows":
        raise StagingLoadError("structured loader accepts only datarows workbooks")
    if source.logical_table_id != table.logical_table_id:
        raise StagingLoadError("source logical table does not match staging table")
    if data_profile_table.get("logical_table_id") != table.logical_table_id:
        raise StagingLoadError("data profile logical table does not match staging table")
    _validate_rule_alignment(table, rules)
    if not verify_sources_unchanged((source,)):
        raise StagingLoadError("source workbook changed before staging load")

    expected_headers = [column.qualified_name for column in table.columns]
    sheet_name = str(data_profile_table["selected_data_sheet"])
    header_row = int(data_profile_table["header_row"])
    last_data_row = int(data_profile_table["last_data_row"])
    expected_row_count = int(data_profile_table["row_count"])
    workbook = _load_workbook(source.path)
    observed_rows = 0
    source_row_ids: set[str] = set()
    try:
        if sheet_name not in workbook.sheetnames:
            raise StagingLoadError(
                f"source data sheet {sheet_name!r} is missing for {table.logical_table_id}"
            )
        worksheet = workbook[sheet_name]
        header_values = next(
            worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True)
        )
        if _qualified_headers(header_values) != expected_headers:
            raise StagingLoadError(
                f"source header differs from staging schema for {table.logical_table_id}"
            )

        for excel_row, cells in enumerate(
            worksheet.iter_rows(
                min_row=header_row + 1,
                max_row=last_data_row,
                max_col=len(table.columns),
            ),
            start=header_row + 1,
        ):
            # openpyxl returns None for a physically present empty inlineStr.
            # Restore that exact state so it remains EMPTY rather than MISSING.
            values = tuple(
                "" if cell.value is None and cell.data_type == "inlineStr" else cell.value
                for cell in cells
            )
            if not any(not is_blank_value(value) for value in values):
                continue
            lineage = build_source_lineage(
                source_table=table.logical_table_id,
                source_file_sha256=source.sha256,
                source_sheet=sheet_name,
                source_excel_row=excel_row,
            )
            if lineage.source_row_id in source_row_ids:
                raise StagingLoadError(
                    f"source_row_id collision for {table.logical_table_id}"
                )
            source_row_ids.add(lineage.source_row_id)
            observed_rows += 1

            raw_record: dict[str, Mapping[str, Any]] = {}
            parse_statuses: dict[str, str] = {}
            typed_values: list[SourceScalar] = []
            for raw_value, rule in zip(values, rules, strict=True):
                qualified_name = str(rule.qualified_name)
                result = normalize_value(raw_value, rule)
                raw_record[qualified_name] = _raw_envelope(raw_value)
                parse_statuses[qualified_name] = result.parse_status.value
                typed_values.append(result.normalized_value)
            yield StagedRow(
                lineage=lineage,
                raw_record=raw_record,
                parse_statuses=parse_statuses,
                typed_values=tuple(typed_values),
            )
    finally:
        workbook.close()

    if observed_rows != expected_row_count:
        raise StagingLoadError(
            f"source row count mismatch for {table.logical_table_id}: "
            f"expected {expected_row_count}, observed {observed_rows}"
        )
    if len(source_row_ids) != observed_rows:
        raise StagingLoadError(
            f"source_row_id uniqueness failure for {table.logical_table_id}"
        )


def render_insert_sql(table: StagingTableSpec) -> str:
    """Return a parameter-only insert statement for a validated table spec."""

    metadata = (
        "source_row_id",
        "logical_table_id",
        "source_excel_row",
        "source_file_sha256",
        "source_sheet",
        "raw_record",
        "parse_statuses",
    )
    column_names = (*metadata, *(column.db_column_name for column in table.columns))
    identifiers = ", ".join(quote_identifier(name) for name in column_names)
    placeholders = ["%s"] * len(metadata)
    placeholders[5] = "%s::jsonb"
    placeholders[6] = "%s::jsonb"
    # Explicit casts bind each value to the declared staging type and avoid
    # relying on driver inference.  Type strings come only from the fixed
    # TargetType mapping in schema.py.
    placeholders.extend(
        f"%s::{column.postgres_type}" for column in table.columns
    )
    return (
        f"INSERT INTO {table.qualified_sql_name} ({identifiers}) VALUES "
        f"({', '.join(placeholders)})"
    )


def _fetch_counts(cursor: CursorLike, table: StagingTableSpec) -> tuple[int, int]:
    cursor.execute(
        "SELECT count(*) AS row_count, "
        "count(DISTINCT \"source_row_id\") AS distinct_source_row_id_count "
        f"FROM {table.qualified_sql_name}"
    )
    row = cursor.fetchone()
    if isinstance(row, Mapping):
        values = (
            row.get("row_count") if "row_count" in row else row.get("count"),
            row.get("distinct_source_row_id_count"),
        )
    else:
        try:
            values = row[0], row[1]
        except (TypeError, IndexError) as error:
            raise StagingLoadError("database count query returned an invalid row") from error
    try:
        return int(values[0]), int(values[1])
    except (TypeError, ValueError) as error:
        raise StagingLoadError("database count query returned non-integer values") from error


def validate_loaded_table(
    connection: ConnectionLike,
    table: StagingTableSpec,
    *,
    expected_row_count: int,
) -> tuple[int, int]:
    """Verify total rows and technical-key uniqueness for one loaded table."""

    with connection.cursor() as cursor:
        row_count, distinct_count = _fetch_counts(cursor, table)
    if row_count != expected_row_count or distinct_count != expected_row_count:
        raise StagingLoadError(
            f"database validation failed for {table.logical_table_id}: "
            f"expected={expected_row_count}, rows={row_count}, "
            f"distinct_source_row_ids={distinct_count}"
        )
    return row_count, distinct_count


def expected_catalog_columns(
    table: StagingTableSpec,
) -> tuple[tuple[str, str, str, int], ...]:
    """Return the exact information_schema column projection for a table."""

    fixed = (
        ("source_row_id", "text", "NO"),
        ("logical_table_id", "text", "NO"),
        ("source_excel_row", "bigint", "NO"),
        ("source_file_sha256", "text", "NO"),
        ("source_sheet", "text", "NO"),
        ("raw_record", "jsonb", "NO"),
        ("parse_statuses", "jsonb", "NO"),
    )
    dynamic = tuple(
        (column.db_column_name, column.postgres_type, "YES")
        for column in table.columns
    )
    return tuple(
        (name, data_type, nullable, position)
        for position, (name, data_type, nullable) in enumerate(
            (*fixed, *dynamic), start=1
        )
    )


def _catalog_row(row: Any) -> tuple[str, str, str, int]:
    if isinstance(row, Mapping):
        values = (
            row.get("column_name"),
            row.get("data_type"),
            row.get("is_nullable"),
            row.get("ordinal_position"),
        )
    else:
        try:
            values = row[0], row[1], row[2], row[3]
        except (TypeError, IndexError) as error:
            raise StagingLoadError("catalog column query returned an invalid row") from error
    try:
        return str(values[0]), str(values[1]), str(values[2]), int(values[3])
    except (TypeError, ValueError) as error:
        raise StagingLoadError("catalog column query returned invalid values") from error


def _primary_key_column(row: Any) -> str:
    if isinstance(row, Mapping):
        value = row.get("column_name")
    else:
        try:
            value = row[0]
        except (TypeError, IndexError) as error:
            raise StagingLoadError("catalog primary-key query returned an invalid row") from error
    return str(value)


def _unique_constraint_part(row: Any) -> tuple[str, str, int]:
    if isinstance(row, Mapping):
        values = (
            row.get("constraint_name"),
            row.get("column_name"),
            row.get("ordinal_position"),
        )
    else:
        try:
            values = row[0], row[1], row[2]
        except (TypeError, IndexError) as error:
            raise StagingLoadError(
                "catalog unique-constraint query returned an invalid row"
            ) from error
    try:
        return str(values[0]), str(values[1]), int(values[2])
    except (TypeError, ValueError) as error:
        raise StagingLoadError(
            "catalog unique-constraint query returned invalid values"
        ) from error


def _check_clause(row: Any) -> str:
    if isinstance(row, Mapping):
        value = row.get("check_clause")
    else:
        try:
            value = row[0]
        except (TypeError, IndexError) as error:
            raise StagingLoadError(
                "catalog check-constraint query returned an invalid row"
            ) from error
    if not isinstance(value, str):
        raise StagingLoadError("catalog check constraint is not text")
    return value


def _compact_check_clause(clause: str) -> str:
    compact = re.sub(r'[\s"()]', "", clause).casefold()
    return compact.replace("::text", "").replace("::bigint", "")


def _validate_catalog_with_cursor(
    cursor: CursorLike,
    tables: Iterable[StagingTableSpec],
) -> None:
    for table in tables:
        parameters = (table.db_schema_name, table.db_table_name)
        cursor.execute(_CATALOG_COLUMNS_SQL, parameters)
        observed_columns = tuple(_catalog_row(row) for row in cursor.fetchall())
        expected_columns = expected_catalog_columns(table)
        if observed_columns != expected_columns:
            raise StagingLoadError(
                f"existing PostgreSQL table shape is incompatible for "
                f"{table.logical_table_id}"
            )
        cursor.execute(_CATALOG_PRIMARY_KEY_SQL, parameters)
        primary_key = tuple(_primary_key_column(row) for row in cursor.fetchall())
        if primary_key != ("source_row_id",):
            raise StagingLoadError(
                f"existing PostgreSQL primary key is incompatible for "
                f"{table.logical_table_id}: {primary_key}"
            )
        cursor.execute(_CATALOG_UNIQUE_SQL, parameters)
        unique_parts: dict[str, list[tuple[int, str]]] = {}
        for row in cursor.fetchall():
            constraint_name, column_name, position = _unique_constraint_part(row)
            unique_parts.setdefault(constraint_name, []).append((position, column_name))
        observed_unique_columns = {
            tuple(name for _, name in sorted(parts))
            for parts in unique_parts.values()
        }
        required_unique = (
            "source_file_sha256",
            "source_sheet",
            "source_excel_row",
        )
        if required_unique not in observed_unique_columns:
            raise StagingLoadError(
                f"existing PostgreSQL unique lineage constraint is incompatible for "
                f"{table.logical_table_id}"
            )
        cursor.execute(_CATALOG_CHECKS_SQL, parameters)
        checks = tuple(
            _compact_check_clause(_check_clause(row)) for row in cursor.fetchall()
        )
        logical_literal = repr(table.logical_table_id.casefold())
        has_logical_check = any(
            "logical_table_id=" in clause and logical_literal in clause
            for clause in checks
        )
        has_row_check = any("source_excel_row>0" in clause for clause in checks)
        if not has_logical_check or not has_row_check:
            raise StagingLoadError(
                f"existing PostgreSQL check constraints are incompatible for "
                f"{table.logical_table_id}"
            )


def validate_staging_catalog(
    connection: ConnectionLike,
    schema: StagingSchemaSpec,
) -> None:
    """Reject an old/incompatible IF-NOT-EXISTS table before any data write."""

    with connection.cursor() as cursor:
        _validate_catalog_with_cursor(cursor, schema.tables)


def ensure_staging_schema(
    connection: ConnectionLike,
    schema: StagingSchemaSpec,
    *,
    commit: bool = True,
) -> None:
    """Create the fixed schema and its tables transactionally."""

    try:
        with connection.cursor() as cursor:
            cursor.execute(render_create_schema_sql(schema))
            for table in schema.tables:
                cursor.execute(render_create_table_sql(table))
            _validate_catalog_with_cursor(cursor, schema.tables)
        if commit:
            connection.commit()
    except Exception:
        connection.rollback()
        raise


def _load_table_in_transaction(
    connection: ConnectionLike,
    *,
    source: SourceWorkbook,
    table: StagingTableSpec,
    rules: Sequence[TypedRule],
    data_profile_table: Mapping[str, Any],
    batch_size: int,
    replace: bool,
) -> TableLoadStats:
    if isinstance(batch_size, bool) or not isinstance(batch_size, int) or batch_size < 1:
        raise ValueError("batch_size must be a positive integer")
    if not verify_sources_unchanged((source,)):
        raise StagingLoadError("source workbook changed before staging load")
    expected_row_count = int(data_profile_table["row_count"])
    inserted = 0
    batch_count = 0
    status_counts: Counter[str] = Counter()
    insert_sql = render_insert_sql(table)
    with connection.cursor() as cursor:
        if replace:
            cursor.execute(f"DELETE FROM {table.qualified_sql_name}")
        batch: list[tuple[Any, ...]] = []
        for staged_row in iter_typed_rows(
            source=source,
            table=table,
            rules=rules,
            data_profile_table=data_profile_table,
        ):
            batch.append(staged_row.parameters(table))
            status_counts.update(staged_row.parse_statuses.values())
            if len(batch) >= batch_size:
                cursor.executemany(insert_sql, batch)
                inserted += len(batch)
                batch_count += 1
                batch = []
        if batch:
            cursor.executemany(insert_sql, batch)
            inserted += len(batch)
            batch_count += 1
        if inserted != expected_row_count:
            raise StagingLoadError(
                f"insert count mismatch for {table.logical_table_id}: "
                f"expected {expected_row_count}, inserted {inserted}"
            )
        row_count, distinct_count = _fetch_counts(cursor, table)
        if row_count != expected_row_count or distinct_count != expected_row_count:
            raise StagingLoadError(
                f"database validation failed for {table.logical_table_id}: "
                f"expected={expected_row_count}, rows={row_count}, "
                f"distinct_source_row_ids={distinct_count}"
            )
    unchanged = verify_sources_unchanged((source,))
    if not unchanged:
        raise StagingLoadError("source workbook changed during staging load")
    return TableLoadStats(
        logical_table_id=table.logical_table_id,
        source_row_count=expected_row_count,
        inserted_row_count=inserted,
        database_row_count=row_count,
        distinct_source_row_id_count=distinct_count,
        batch_count=batch_count,
        parse_status_counts=dict(sorted(status_counts.items())),
        source_unchanged=unchanged,
    )


def load_table(
    connection: ConnectionLike,
    *,
    source: SourceWorkbook,
    table: StagingTableSpec,
    rules: Sequence[TypedRule],
    data_profile_table: Mapping[str, Any],
    batch_size: int = 1_000,
    replace: bool = True,
    create_table: bool = True,
) -> TableLoadStats:
    """Atomically replace/load/validate one staging table."""

    try:
        if create_table:
            with connection.cursor() as cursor:
                cursor.execute(
                    f"CREATE SCHEMA IF NOT EXISTS {quote_identifier(table.db_schema_name)}"
                )
                cursor.execute(render_create_table_sql(table))
                _validate_catalog_with_cursor(cursor, (table,))
        else:
            with connection.cursor() as cursor:
                _validate_catalog_with_cursor(cursor, (table,))
        stats = _load_table_in_transaction(
            connection,
            source=source,
            table=table,
            rules=rules,
            data_profile_table=data_profile_table,
            batch_size=batch_size,
            replace=replace,
        )
        connection.commit()
        return stats
    except Exception:
        connection.rollback()
        raise


def _index_data_profile(data_profile: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    tables = data_profile.get("tables")
    if not isinstance(tables, list):
        raise StagingLoadError("data profile must contain a tables list")
    indexed: dict[str, Mapping[str, Any]] = {}
    for table in tables:
        if not isinstance(table, Mapping):
            raise StagingLoadError("data profile table must be an object")
        table_id = str(table.get("logical_table_id", ""))
        if not table_id or table_id in indexed:
            raise StagingLoadError(f"invalid or duplicate data profile table: {table_id!r}")
        indexed[table_id] = table
    return indexed


def load_all(
    connection: ConnectionLike,
    *,
    sources: Iterable[SourceWorkbook],
    schema: StagingSchemaSpec,
    rules: Sequence[TypedRule],
    data_profile: Mapping[str, Any],
    batch_size: int = 1_000,
    replace: bool = True,
    create_tables: bool = True,
) -> LoadStats:
    """Atomically load every table; one failure rolls back the whole G2-B load."""

    datarows_sources = [source for source in sources if source.source_kind == "datarows"]
    data_sources = {source.logical_table_id: source for source in datarows_sources}
    if len(data_sources) != len(datarows_sources):
        raise StagingLoadError("duplicate datarows source identity")
    if set(data_sources) != {table.logical_table_id for table in schema.tables}:
        raise StagingLoadError("datarows source identities do not match staging schema")
    profile_tables = _index_data_profile(data_profile)
    if set(profile_tables) != set(data_sources):
        raise StagingLoadError("data profile identities do not match staging sources")
    rules_by_table: dict[str, list[TypedRule]] = {
        table.logical_table_id: [] for table in schema.tables
    }
    for rule in rules:
        if rule.logical_table_id not in rules_by_table:
            raise StagingLoadError(f"unexpected typed-rule table: {rule.logical_table_id}")
        rules_by_table[rule.logical_table_id].append(rule)

    if not all(verify_sources_unchanged((source,)) for source in data_sources.values()):
        raise StagingLoadError("a source workbook changed before all-table load")

    stats: list[TableLoadStats] = []
    try:
        if create_tables:
            with connection.cursor() as cursor:
                cursor.execute(render_create_schema_sql(schema))
                for table in schema.tables:
                    cursor.execute(render_create_table_sql(table))
                _validate_catalog_with_cursor(cursor, schema.tables)
        else:
            with connection.cursor() as cursor:
                _validate_catalog_with_cursor(cursor, schema.tables)
        for table in schema.tables:
            stats.append(
                _load_table_in_transaction(
                    connection,
                    source=data_sources[table.logical_table_id],
                    table=table,
                    rules=rules_by_table[table.logical_table_id],
                    data_profile_table=profile_tables[table.logical_table_id],
                    batch_size=batch_size,
                    replace=replace,
                )
            )
        if not all(verify_sources_unchanged((source,)) for source in data_sources.values()):
            raise StagingLoadError("a source workbook changed during all-table load")
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    return LoadStats(tables=tuple(stats), committed=True)


def current_source_fingerprint(source: SourceWorkbook) -> dict[str, Any]:
    """Return a read-only diagnostic fingerprint without exposing source rows."""

    stat = source.path.stat()
    return {
        "logical_table_id": source.logical_table_id,
        "physical_filename": source.physical_filename,
        "size_bytes": stat.st_size,
        "modified_time_ns": stat.st_mtime_ns,
        "sha256": sha256_file(source.path),
    }


__all__ = [
    "ConnectionLike",
    "CursorLike",
    "LoadStats",
    "StagedRow",
    "StagingLoadError",
    "TableLoadStats",
    "current_source_fingerprint",
    "ensure_staging_schema",
    "expected_catalog_columns",
    "iter_typed_rows",
    "load_all",
    "load_table",
    "render_insert_sql",
    "validate_loaded_table",
    "validate_staging_catalog",
]
