"""Schema extraction and streaming datarows profiling with openpyxl."""

from __future__ import annotations

import tempfile
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from .discovery import SourceWorkbook
from .statistics import (
    ColumnAccumulator,
    DistinctDigestStore,
    composite_digest,
    is_blank_value,
    value_type,
)
from .workbook_metadata import WorkbookMetadata


class WorkbookProfileError(ValueError):
    """Raised when the true schema or datarows header cannot be identified."""


_SCHEMA_HEADER_ALIASES: dict[str, set[str]] = {
    "column_name": {"컬럼명", "column", "columnname", "field", "fieldname"},
    "key_declaration": {"pk/fk", "pkfk", "key", "keytype", "키"},
    "declared_type": {
    "컬럼타입",
    "datatype",
    "columntype",
    "type",
    "자료형",
    "데이터타입",
},

"display_name": {
    "컬럼한글명",
    "한글명",
    "description",
    "displayname",
    "컬럼코멘트",
},
    "declared_nullable": {"nullable", "null여부", "널허용", "필수여부"},
    "declared_unit": {"단위", "unit", "currency", "통화"},
    "example": {"컬럼값예시", "값예시", "example", "sample", "samplevalue"},
}


def _normalise_header(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).strip().casefold()
    return "".join(character for character in text if not character.isspace() and character != "_")


def _schema_header_kind(value: Any) -> str | None:
    normalised = _normalise_header(value)
    for kind, aliases in _SCHEMA_HEADER_ALIASES.items():
        if normalised in aliases:
            return kind
    return None


def _cell_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def qualify_column_names(names: Iterable[str | None]) -> list[dict[str, Any]]:
    """Attach a stable occurrence number so duplicate headers are never collapsed."""

    occurrences: Counter[str | None] = Counter()
    qualified: list[dict[str, Any]] = []
    for position, name in enumerate(names, start=1):
        occurrences[name] += 1
        occurrence = occurrences[name]
        qualified.append(
            {
                "position": position,
                "column_name": name,
                "occurrence": occurrence,
                "qualified_name": f"{name if name is not None else '<blank>'}#{occurrence}",
            }
        )
    return qualified


def duplicate_column_names(columns: list[dict[str, Any]]) -> list[dict[str, Any]]:
    positions: dict[str | None, list[int]] = {}
    for column in columns:
        positions.setdefault(column["column_name"], []).append(column["position"])
    return [
        {
            "column_name": name,
            "positions": column_positions,
            "occurrence_count": len(column_positions),
        }
        for name, column_positions in positions.items()
        if len(column_positions) > 1
    ]


def _load_workbook(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "G1 profiling requires the optional data dependency: "
            "python -m pip install -e \".[data]\""
        ) from error
    return load_workbook(path, read_only=True, data_only=False, keep_links=False)


def _schema_header_candidates(workbook: Any) -> list[tuple[int, int, int, str, dict[int, str]]]:
    candidates: list[tuple[int, int, int, str, dict[int, str]]] = []
    for sheet_order, worksheet in enumerate(workbook.worksheets):
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=50, values_only=True), start=1
        ):
            field_by_position = {
                position: kind
                for position, value in enumerate(row)
                if (kind := _schema_header_kind(value)) is not None
            }
            kinds = set(field_by_position.values())
            if "column_name" in kinds:
                candidates.append(
                    (len(kinds), -sheet_order, -row_number, worksheet.title, field_by_position)
                )
    return candidates


def _is_declared_pk(value: str | None) -> bool:
    if value is None:
        return False
    tokens = {
        token
        for token in value.upper().replace("PRIMARY KEY", "PK").replace("/", " ").split()
    }
    return "PK" in tokens


def extract_schema_profile(
    source: SourceWorkbook,
    metadata: WorkbookMetadata,
) -> dict[str, Any]:
    """Extract only explicitly declared schema facts; example values are omitted."""

    workbook = _load_workbook(source.path)
    try:
        candidates = _schema_header_candidates(workbook)
        if not candidates:
            raise WorkbookProfileError(
                f"schema header containing a column-name field was not found: {source.path}"
            )
        _, _, negative_row, sheet_name, field_by_position = max(candidates)
        header_row = -negative_row
        worksheet = workbook[sheet_name]
        raw_header = next(
            worksheet.iter_rows(
                min_row=header_row,
                max_row=header_row,
                values_only=True,
            )
        )
        schema_headers = [_cell_text(value) for value in raw_header]

        raw_columns: list[dict[str, Any]] = []
        name_position = next(
            position
            for position, kind in field_by_position.items()
            if kind == "column_name"
        )
        for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            name = _cell_text(row[name_position] if name_position < len(row) else None)
            if name is None:
                continue
            declared: dict[str, Any] = {"column_name": name}
            for position, kind in field_by_position.items():
                if kind in {"column_name", "example"}:
                    continue
                declared[kind] = _cell_text(row[position] if position < len(row) else None)
            raw_columns.append(declared)

        qualified = qualify_column_names(column["column_name"] for column in raw_columns)
        columns: list[dict[str, Any]] = []
        for identity, declaration in zip(qualified, raw_columns, strict=True):
            columns.append(
                {
                    **identity,
                    "key_declaration": declaration.get("key_declaration"),
                    "declared_primary_key": _is_declared_pk(
                        declaration.get("key_declaration")
                    ),
                    "declared_type": declaration.get("declared_type"),
                    "display_name": declaration.get("display_name"),
                    "declared_nullable": declaration.get("declared_nullable"),
                    "declared_unit": declaration.get("declared_unit"),
                }
            )

        selected_metadata = metadata.sheet(sheet_name)
        return {
            "logical_table_id": source.logical_table_id,
            "physical_filename": source.physical_filename,
            "sheet_names": metadata.sheet_names,
            "selected_schema_sheet": sheet_name,
            "header_row": header_row,
            "schema_metadata_column_count": sum(
                header is not None for header in schema_headers
            ),
            "schema_headers": schema_headers,
            "schema_definition_count": len(columns),
            "duplicate_column_names": duplicate_column_names(columns),
            "columns": columns,
            "declared_primary_key_columns": [
                column["qualified_name"]
                for column in columns
                if column["declared_primary_key"]
            ],
            "declared_unit_column_count": sum(
                column["declared_unit"] is not None for column in columns
            ),
            "structure": selected_metadata.to_dict(),
        }
    finally:
        workbook.close()


def _trim_trailing_blank_cells(row: tuple[Any, ...]) -> tuple[Any, ...]:
    end = len(row)
    while end and row[end - 1] is None:
        end -= 1
    return row[:end]


def _find_data_header(
    workbook: Any,
    schema_columns: list[dict[str, Any]],
) -> tuple[str, int, list[str | None]]:
    expected_names = Counter(column["column_name"] for column in schema_columns)
    candidates: list[tuple[int, int, int, int, str, tuple[Any, ...]]] = []
    for sheet_order, worksheet in enumerate(workbook.worksheets):
        for row_number, raw_row in enumerate(
            worksheet.iter_rows(min_row=1, max_row=50, values_only=True), start=1
        ):
            row = _trim_trailing_blank_cells(raw_row)
            names = [_cell_text(value) for value in row]
            nonblank_names = [name for name in names if name is not None]
            if not nonblank_names:
                continue
            overlap = sum((Counter(nonblank_names) & expected_names).values())
            unique_count = len(set(nonblank_names))
            candidates.append(
                (
                    overlap,
                    len(nonblank_names),
                    unique_count,
                    -sheet_order * 1000 - row_number,
                    worksheet.title,
                    row,
                )
            )
    if not candidates:
        raise WorkbookProfileError("no nonblank candidate header row was found")
    overlap, _, _, order_score, sheet_name, raw_header = max(candidates)
    if overlap == 0:
        raise WorkbookProfileError(
            "datarows header has no exact column overlap with the declared schema"
        )
    row_number = -(order_score % 1000) if order_score < 0 else 0
    # The compact order score above is only a tie-breaker. Locate the exact row
    # again so sheet order cannot affect the reported one-based row number.
    worksheet = workbook[sheet_name]
    for candidate_row_number, candidate_row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=50, values_only=True), start=1
    ):
        if _trim_trailing_blank_cells(candidate_row) == raw_header:
            row_number = candidate_row_number
            break
    return sheet_name, row_number, [_cell_text(value) for value in raw_header]


def _data_column_lookup(
    data_columns: list[dict[str, Any]],
) -> dict[str, dict[str, Any]]:
    return {column["qualified_name"]: column for column in data_columns}


def profile_data_workbook(
    source: SourceWorkbook,
    schema_profile: dict[str, Any],
    metadata: WorkbookMetadata,
) -> dict[str, Any]:
    """Profile one datarows workbook in a single value pass."""

    workbook = _load_workbook(source.path)
    try:
        sheet_name, header_row, header_names = _find_data_header(
            workbook, schema_profile["columns"]
        )
        worksheet = workbook[sheet_name]
        data_columns = qualify_column_names(header_names)
        schema_by_name = {
            column["qualified_name"]: column for column in schema_profile["columns"]
        }
        for column in data_columns:
            declaration = schema_by_name.get(column["qualified_name"], {})
            column.update(
                {
                    "declared_type": declaration.get("declared_type"),
                    "declared_nullable": declaration.get("declared_nullable"),
                    "declared_unit": declaration.get("declared_unit"),
                    "key_declaration": declaration.get("key_declaration"),
                    "declared_primary_key": declaration.get(
                        "declared_primary_key", False
                    ),
                }
            )

        data_lookup = _data_column_lookup(data_columns)
        declared_key_names = schema_profile["declared_primary_key_columns"]
        declared_key_positions = [
            data_lookup[name]["position"] - 1
            for name in declared_key_names
            if name in data_lookup
        ]
        missing_declared_key_columns = [
            name for name in declared_key_names if name not in data_lookup
        ]
        composite_slot = len(data_columns)

        with tempfile.TemporaryDirectory(prefix="g1-profile-distinct-") as temp_dir:
            with DistinctDigestStore(Path(temp_dir) / "distinct.sqlite3") as distinct_store:
                accumulators = [
                    ColumnAccumulator(slot=index, distinct_store=distinct_store)
                    for index in range(len(data_columns))
                ]
                row_count = 0
                blank_rows_between_header_and_last_data = 0
                pending_blank_rows = 0
                last_data_row = header_row
                rows_with_values_beyond_header = 0
                key_rows_with_any_blank = 0
                key_complete_rows = 0

                for excel_row_number, row in enumerate(
                    worksheet.iter_rows(min_row=header_row + 1), start=header_row + 1
                ):
                    row_has_data = any(not is_blank_value(cell.value) for cell in row)
                    if not row_has_data:
                        pending_blank_rows += 1
                        continue
                    blank_rows_between_header_and_last_data += pending_blank_rows
                    pending_blank_rows = 0
                    row_count += 1
                    last_data_row = excel_row_number

                    if any(
                        not is_blank_value(cell.value) for cell in row[len(data_columns) :]
                    ):
                        rows_with_values_beyond_header += 1

                    for index, accumulator in enumerate(accumulators):
                        if index < len(row):
                            cell = row[index]
                            accumulator.update(
                                cell.value,
                                excel_data_type=getattr(cell, "data_type", None),
                            )
                        else:
                            accumulator.update(None)

                    if declared_key_names and not missing_declared_key_columns:
                        key_values: list[tuple[Any, str]] = []
                        key_has_blank = False
                        for index in declared_key_positions:
                            cell = row[index] if index < len(row) else None
                            value = cell.value if cell is not None else None
                            if is_blank_value(value):
                                key_has_blank = True
                                break
                            key_values.append(
                                (
                                    value,
                                    value_type(
                                        value,
                                        excel_data_type=getattr(cell, "data_type", None),
                                    ),
                                )
                            )
                        if key_has_blank:
                            key_rows_with_any_blank += 1
                        else:
                            key_complete_rows += 1
                            distinct_store.add_digest(
                                composite_slot, composite_digest(key_values)
                            )

                distinct_counts = distinct_store.counts()
                profiled_columns: list[dict[str, Any]] = []
                for index, (identity, accumulator) in enumerate(
                    zip(data_columns, accumulators, strict=True)
                ):
                    profiled_columns.append(
                        {
                            **identity,
                            **accumulator.finalise(
                                row_count=row_count,
                                distinct_count=distinct_counts.get(index, 0),
                            ),
                        }
                    )
                key_distinct_count = distinct_counts.get(composite_slot, 0)

        candidate_identifiers = [
            {
                "qualified_name": column["qualified_name"],
                "column_name": column["column_name"],
                "null_count": column["null_count"],
                "distinct_count": column["distinct_count"],
                "duplicate_count": column["duplicate_count"],
                "basis": "observed_unique_non_null_only",
            }
            for column in profiled_columns
            if row_count > 0
            and column["null_count"] == 0
            and column["distinct_count"] == row_count
            and column["excel_error_count"] == 0
            and column["formula_cell_count"] == 0
        ]

        if not declared_key_names:
            declared_key_analysis: dict[str, Any] = {
                "status": "NOT_DECLARED_IN_SCHEMA",
                "columns": [],
                "interpretation": None,
            }
        elif missing_declared_key_columns:
            declared_key_analysis = {
                "status": "NOT_EVALUABLE",
                "columns": declared_key_names,
                "missing_in_data": missing_declared_key_columns,
                "interpretation": "schema declaration preserved; key semantics not inferred",
            }
        else:
            declared_key_analysis = {
                "status": "SCHEMA_DECLARED",
                "kind": "SINGLE" if len(declared_key_names) == 1 else "COLUMN_SET",
                "columns": declared_key_names,
                "rows_with_any_null_or_blank": key_rows_with_any_blank,
                "complete_rows": key_complete_rows,
                "distinct_complete_combinations": key_distinct_count,
                "duplicate_complete_combinations": max(
                    0, key_complete_rows - key_distinct_count
                ),
                "unique_when_complete": key_complete_rows == key_distinct_count,
                "interpretation": (
                    "multiple PK markers are evaluated together; composite-key "
                    "semantics are not asserted"
                    if len(declared_key_names) > 1
                    else "single schema PK marker; observed uniqueness is reported separately"
                ),
            }

        selected_metadata = metadata.sheet(sheet_name)
        structural_max_row = max(
            selected_metadata.xml_last_row,
            selected_metadata.dimension_max_row or 0,
        )
        return {
            "logical_table_id": source.logical_table_id,
            "physical_filename": source.physical_filename,
            "sheet_names": metadata.sheet_names,
            "selected_data_sheet": sheet_name,
            "header_row": header_row,
            "column_count": len(profiled_columns),
            "row_count": row_count,
            "last_data_row": last_data_row,
            "blank_trailing_rows": max(0, structural_max_row - last_data_row),
            "blank_rows_between_header_and_last_data": blank_rows_between_header_and_last_data,
            "rows_with_values_beyond_header": rows_with_values_beyond_header,
            "duplicate_column_names": duplicate_column_names(profiled_columns),
            "declared_key_analysis": declared_key_analysis,
            "candidate_identifiers": candidate_identifiers,
            "columns": profiled_columns,
            "structure": selected_metadata.to_dict(),
        }
    finally:
        workbook.close()
