"""Streaming official-workbook scan for aggregate G2-A evidence.

The scanner never emits source values. It keeps only parse counters, technical
lineage identifiers, and SHA-256 digests needed for exact trim-collision
counts. Raw cell values live only for the duration of processing one row.
"""

from __future__ import annotations

import hashlib
import struct
from collections import Counter, defaultdict
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable, Mapping

from pipeline.profiling.discovery import SourceWorkbook
from pipeline.profiling.statistics import is_blank_value

from .lineage import SOURCE_ROW_ID_VERSION, derive_source_row_id
from .models import CodeStatus, MappingStatus, TypedRule, UnitStatus
from .normalizer import normalize_value
from .type_rules import build_typed_rule
from .whitespace import TrimStatus


class TypedStagingScanError(ValueError):
    """Raised when G1 evidence and a physical workbook do not agree."""


def _identity(table_id: str, qualified_name: str) -> tuple[str, str]:
    return table_id, qualified_name


def _table_index(document: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    tables = document.get("tables")
    if not isinstance(tables, list):
        raise TypedStagingScanError("profile document must contain a tables list")
    indexed: dict[str, Mapping[str, Any]] = {}
    for table in tables:
        if not isinstance(table, Mapping):
            raise TypedStagingScanError("profile table must be an object")
        table_id = str(table.get("logical_table_id", ""))
        if not table_id or table_id in indexed:
            raise TypedStagingScanError(f"invalid or duplicate table: {table_id!r}")
        indexed[table_id] = table
    return indexed


def _column_index(table: Mapping[str, Any]) -> dict[str, Mapping[str, Any]]:
    columns = table.get("columns")
    if not isinstance(columns, list):
        raise TypedStagingScanError("profile table columns must be a list")
    indexed: dict[str, Mapping[str, Any]] = {}
    for column in columns:
        if not isinstance(column, Mapping):
            raise TypedStagingScanError("profile column must be an object")
        qualified = str(column.get("qualified_name", ""))
        if not qualified or qualified in indexed:
            raise TypedStagingScanError(
                f"invalid or duplicate qualified column: {qualified!r}"
            )
        indexed[qualified] = column
    return indexed


def build_rules_from_profiles(
    schema_profile: Mapping[str, Any],
    data_profile: Mapping[str, Any],
) -> list[TypedRule]:
    """Build one conservative typed rule per occurrence-aware source column."""

    schema_tables = _table_index(schema_profile)
    data_tables = _table_index(data_profile)
    if list(schema_tables) != list(data_tables):
        raise TypedStagingScanError("schema/data profile table order differs")

    rules: list[TypedRule] = []
    for table_id, data_table in data_tables.items():
        schema_columns = _column_index(schema_tables[table_id])
        data_columns = data_table.get("columns", [])
        if list(schema_columns) != [column["qualified_name"] for column in data_columns]:
            raise TypedStagingScanError(
                f"schema/data column order differs for {table_id}"
            )

        for data_column in data_columns:
            qualified = str(data_column["qualified_name"])
            schema_column = schema_columns[qualified]
            display_name = schema_column.get("display_name")
            declared_unit = schema_column.get("declared_unit")
            code_candidate = bool(data_column.get("enum_code_candidate")) or bool(
                data_column.get("boolean_code_candidate")
            )
            code_shape = data_column.get("code_like_string")
            if isinstance(code_shape, Mapping):
                code_candidate = code_candidate or bool(code_shape.get("candidate"))

            rules.append(
                build_typed_rule(
                    logical_table_id=table_id,
                    source_column=str(data_column["column_name"]),
                    occurrence=int(data_column.get("occurrence", 1)),
                    qualified_name=qualified,
                    declared_type=str(schema_column["declared_type"]),
                    observed_types=tuple(data_column.get("observed_types", [])),
                    unit_status=(
                        UnitStatus.CONFIRMED
                        if declared_unit not in (None, "")
                        else UnitStatus.UNMAPPED
                    ),
                    code_status=(
                        CodeStatus.CANDIDATE
                        if code_candidate
                        else CodeStatus.UNMAPPED
                    ),
                    mapping_status=(
                        MappingStatus.CANDIDATE
                        if display_name not in (None, "")
                        else MappingStatus.UNMAPPED
                    ),
                )
            )
    return rules


def _typed_bytes(value: Any) -> bytes:
    if isinstance(value, bool):
        label, encoded = b"boolean", b"1" if value else b"0"
    elif isinstance(value, datetime):
        label, encoded = b"datetime", value.isoformat().encode("utf-8")
    elif isinstance(value, date):
        label, encoded = b"date", value.isoformat().encode("utf-8")
    elif isinstance(value, time):
        label, encoded = b"time", value.isoformat().encode("utf-8")
    elif isinstance(value, Decimal):
        label, encoded = b"decimal", str(value).encode("ascii")
    elif isinstance(value, int):
        label, encoded = b"integer", str(value).encode("ascii")
    elif isinstance(value, float):
        label, encoded = b"number", value.hex().encode("ascii")
    elif isinstance(value, str):
        label, encoded = b"string", value.encode("utf-8", errors="surrogatepass")
    else:
        label = type(value).__name__.encode("ascii", errors="replace")
        encoded = repr(value).encode("utf-8", errors="surrogatepass")
    return label + b":" + encoded


def _typed_digest(value: Any) -> bytes:
    return hashlib.sha256(_typed_bytes(value)).digest()


def _tuple_digest(values: Iterable[Any]) -> bytes:
    digest = hashlib.sha256()
    for value in values:
        encoded = _typed_bytes(value)
        digest.update(struct.pack(">I", len(encoded)))
        digest.update(encoded)
    return digest.digest()


class _WhitespaceDigestAccumulator:
    """Exact-on-digests trim analysis without retaining a source value."""

    def __init__(self) -> None:
        self.total_count = 0
        self.missing_count = 0
        self.empty_string_count = 0
        self.string_count = 0
        self.trimmed_value_count = 0
        self.whitespace_only_count = 0
        self.before_digests: set[bytes] = set()
        self.after_groups: dict[bytes, set[bytes]] = defaultdict(set)

    def update(self, value: Any) -> None:
        self.total_count += 1
        if value is None:
            self.missing_count += 1
            return

        before_digest = _typed_digest(value)
        after_value = value
        if isinstance(value, str):
            self.string_count += 1
            if value == "":
                self.empty_string_count += 1
            elif value.isspace():
                self.whitespace_only_count += 1
            after_value = value.strip()
            if after_value != value:
                self.trimmed_value_count += 1
        after_digest = _typed_digest(after_value)
        self.before_digests.add(before_digest)
        self.after_groups[after_digest].add(before_digest)

    def finalise(self, *, identifier: bool) -> dict[str, Any]:
        present_count = self.total_count - self.missing_count
        null_class_count = (
            self.missing_count
            + self.empty_string_count
            + self.whitespace_only_count
        )
        before = len(self.before_digests)
        after = len(self.after_groups)
        collision_groups = sum(
            1 for source_digests in self.after_groups.values() if len(source_digests) > 1
        )
        unique_before = (
            null_class_count == 0
            and present_count > 0
            and before == self.total_count
        ) if identifier else None
        unique_after = (
            null_class_count == 0
            and present_count > 0
            and after == self.total_count
        ) if identifier else None
        uniqueness_changed = (
            unique_before != unique_after if identifier else None
        )
        conflict = collision_groups > 0 or uniqueness_changed is True
        if conflict:
            trim_status = TrimStatus.CONFLICT
        elif self.trimmed_value_count:
            trim_status = TrimStatus.SAFE
        else:
            trim_status = TrimStatus.NOT_APPLICABLE
        return {
            "total_count": self.total_count,
            "missing_count": self.missing_count,
            "empty_string_count": self.empty_string_count,
            "null_class_count": null_class_count,
            "present_count": present_count,
            "string_count": self.string_count,
            "trimmed_value_count": self.trimmed_value_count,
            "whitespace_only_count": self.whitespace_only_count,
            "distinct_count_before_trim": before,
            "distinct_count_after_trim": after,
            "collapsed_distinct_count": before - after,
            "collision_group_count": collision_groups,
            "trim_collision_count": collision_groups,
            "trim_collision_detected": collision_groups > 0,
            "trim_status": trim_status.value,
            "automatic_trim_applied": False,
            "identifier_assessed": identifier,
            "identifier_unique_before_trim": unique_before,
            "identifier_unique_after_trim": unique_after,
            "identifier_uniqueness_changed": uniqueness_changed,
        }


def _qualified_headers(values: tuple[Any, ...]) -> list[str]:
    end = len(values)
    while end and values[end - 1] is None:
        end -= 1
    counts: Counter[str] = Counter()
    qualified: list[str] = []
    for value in values[:end]:
        if value is None or str(value).strip() == "":
            name = "<blank>"
        else:
            name = str(value).strip()
        counts[name] += 1
        qualified.append(f"{name}#{counts[name]}")
    return qualified


def _load_workbook(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - installation guidance
        raise RuntimeError(
            "G2-A scan requires: python -m pip install -e \".[data]\""
        ) from error
    return load_workbook(path, read_only=True, data_only=False, keep_links=False)


def _shortcut_whitespace(
    data_column: Mapping[str, Any],
    *,
    identifier: bool,
) -> dict[str, Any]:
    total_count = int(data_column["row_count"])
    missing_count = int(data_column.get("missing_cell_count", 0))
    empty_string_count = int(data_column.get("empty_string_count", 0))
    whitespace_only_count = int(data_column.get("whitespace_only_count", 0))
    null_class_count = missing_count + empty_string_count + whitespace_only_count
    present_count = total_count - missing_count
    # G1 distinct counts intentionally exclude empty/null-class strings. With
    # no trim-affected strings, only an observed exact empty string adds one
    # raw distinct value for this G2-A raw-preserving analysis.
    distinct_count = int(data_column["distinct_count"]) + (
        1 if empty_string_count > 0 else 0
    )
    unique = (
        null_class_count == 0
        and total_count > 0
        and distinct_count == total_count
    )
    return {
        "total_count": total_count,
        "missing_count": missing_count,
        "empty_string_count": empty_string_count,
        "null_class_count": null_class_count,
        "present_count": present_count,
        "string_count": int(
            data_column.get("observed_type_counts", {}).get("string", 0)
        ) + empty_string_count,
        "trimmed_value_count": 0,
        "whitespace_only_count": whitespace_only_count,
        "distinct_count_before_trim": distinct_count,
        "distinct_count_after_trim": distinct_count,
        "collapsed_distinct_count": 0,
        "collision_group_count": 0,
        "trim_collision_count": 0,
        "trim_collision_detected": False,
        "trim_status": TrimStatus.NOT_APPLICABLE.value,
        "automatic_trim_applied": False,
        "identifier_assessed": identifier,
        "identifier_unique_before_trim": unique if identifier else None,
        "identifier_unique_after_trim": unique if identifier else None,
        "identifier_uniqueness_changed": False if identifier else None,
    }


def scan_workbooks(
    *,
    sources: tuple[SourceWorkbook, ...],
    data_profile: Mapping[str, Any],
    rules: list[TypedRule],
    parse_error_sample_limit: int = 20,
) -> dict[str, Any]:
    """Scan datarows once and return aggregate parse, trim, and lineage evidence."""

    data_sources = {
        source.logical_table_id: source
        for source in sources
        if source.source_kind == "datarows"
    }
    tables = _table_index(data_profile)
    rule_index = {
        _identity(rule.logical_table_id, str(rule.qualified_name)): rule
        for rule in rules
    }
    if len(rule_index) != len(rules):
        raise TypedStagingScanError("typed rules contain duplicate source identities")

    parse_aggregates: list[dict[str, Any]] = []
    whitespace_aggregates: list[dict[str, Any]] = []
    lineage_tables: list[dict[str, Any]] = []
    key_trim_analyses: list[dict[str, Any]] = []
    total_rows = 0

    for table_id, table in tables.items():
        source = data_sources.get(table_id)
        if source is None:
            raise TypedStagingScanError(f"missing datarows source for {table_id}")
        columns = list(table["columns"])
        qualified_names = [str(column["qualified_name"]) for column in columns]
        table_rules = [rule_index[_identity(table_id, name)] for name in qualified_names]
        candidate_names = {
            str(candidate["qualified_name"])
            for candidate in table.get("candidate_identifiers", [])
        }
        declared_names = set(table.get("declared_key_analysis", {}).get("columns", []))
        identifier_names = candidate_names | declared_names

        status_counters = [Counter() for _ in columns]
        parse_error_samples: list[list[dict[str, Any]]] = [[] for _ in columns]
        whitespace_accumulators: dict[int, _WhitespaceDigestAccumulator] = {
            index: _WhitespaceDigestAccumulator()
            for index, column in enumerate(columns)
            if int(column.get("leading_or_trailing_whitespace_count", 0)) > 0
            or str(column["qualified_name"]) in identifier_names
        }

        declared_positions = [
            qualified_names.index(name) for name in table.get("declared_key_analysis", {}).get("columns", [])
        ]
        key_before: set[bytes] = set()
        key_after: set[bytes] = set()
        key_blank_rows = 0

        workbook = _load_workbook(source.path)
        try:
            sheet_name = str(table["selected_data_sheet"])
            if sheet_name not in workbook.sheetnames:
                raise TypedStagingScanError(
                    f"data sheet {sheet_name!r} missing for {table_id}"
                )
            worksheet = workbook[sheet_name]
            header_row = int(table["header_row"])
            header_values = next(
                worksheet.iter_rows(
                    min_row=header_row,
                    max_row=header_row,
                    values_only=True,
                )
            )
            if _qualified_headers(header_values) != qualified_names:
                raise TypedStagingScanError(
                    f"physical header no longer matches G1 profile for {table_id}"
                )

            observed_rows = 0
            first_excel_row: int | None = None
            last_excel_row: int | None = None
            first_source_row_id: str | None = None
            last_source_row_id: str | None = None
            source_row_ids: set[str] = set()

            for excel_row, cells in enumerate(
                worksheet.iter_rows(
                    min_row=header_row + 1,
                    max_row=int(table["last_data_row"]),
                    max_col=len(columns),
                ),
                start=header_row + 1,
            ):
                # openpyxl represents a physically present empty inline-string
                # cell as ``value=None`` plus ``data_type='inlineStr'``.  G1
                # deliberately distinguishes that source state from a missing
                # cell, so reconstruct only that exact empty-string value here.
                values = tuple(
                    ""
                    if cell.value is None and cell.data_type == "inlineStr"
                    else cell.value
                    for cell in cells
                )
                if not any(not is_blank_value(value) for value in values):
                    continue
                observed_rows += 1
                total_rows += 1
                source_row_id = derive_source_row_id(
                    source_table=table_id,
                    source_file_sha256=source.sha256,
                    source_sheet=sheet_name,
                    source_excel_row=excel_row,
                )
                source_row_ids.add(source_row_id)
                if first_excel_row is None:
                    first_excel_row = excel_row
                    first_source_row_id = source_row_id
                last_excel_row = excel_row
                last_source_row_id = source_row_id

                for index, (raw_value, rule) in enumerate(zip(values, table_rules, strict=True)):
                    result = normalize_value(raw_value, rule)
                    status_counters[index][result.parse_status.value] += 1
                    if (
                        result.parse_status.value == "PARSE_ERROR"
                        and len(parse_error_samples[index]) < parse_error_sample_limit
                    ):
                        parse_error_samples[index].append(
                            {
                                "source_excel_row": excel_row,
                                "source_row_id": source_row_id,
                            }
                        )
                    accumulator = whitespace_accumulators.get(index)
                    if accumulator is not None:
                        accumulator.update(raw_value)

                if declared_positions:
                    raw_parts: list[Any] = []
                    trimmed_parts: list[Any] = []
                    key_blank = False
                    for position in declared_positions:
                        value = values[position]
                        if is_blank_value(value):
                            key_blank = True
                            break
                        raw_parts.append(value)
                        trimmed_parts.append(value.strip() if isinstance(value, str) else value)
                    if key_blank:
                        key_blank_rows += 1
                    else:
                        key_before.add(_tuple_digest(raw_parts))
                        key_after.add(_tuple_digest(trimmed_parts))

            expected_rows = int(table["row_count"])
            if observed_rows != expected_rows:
                raise TypedStagingScanError(
                    f"row count changed for {table_id}: expected {expected_rows}, "
                    f"observed {observed_rows}"
                )
            if len(source_row_ids) != observed_rows:
                raise TypedStagingScanError(
                    f"source_row_id collision detected for {table_id}"
                )
        finally:
            workbook.close()

        for index, (column, rule) in enumerate(zip(columns, table_rules, strict=True)):
            counts = status_counters[index]
            aggregate = {
                "logical_table_id": table_id,
                "source_column": rule.source_column,
                "occurrence": rule.occurrence,
                "qualified_name": rule.qualified_name,
                "row_count": int(table["row_count"]),
                "status_counts": {
                    status: counts.get(status, 0)
                    for status in (
                        "OK",
                        "MISSING",
                        "EMPTY",
                        "WHITESPACE_ONLY",
                        "PARSE_ERROR",
                        "NOT_APPLICABLE",
                    )
                },
                "parse_error_count": counts.get("PARSE_ERROR", 0),
                "parse_error_lineage_samples": parse_error_samples[index],
                "raw_values_included": False,
            }
            if sum(aggregate["status_counts"].values()) != aggregate["row_count"]:
                raise TypedStagingScanError(
                    f"parse status total mismatch for {table_id}.{rule.qualified_name}"
                )
            parse_aggregates.append(aggregate)

            identifier = str(column["qualified_name"]) in identifier_names
            accumulator = whitespace_accumulators.get(index)
            if accumulator is None:
                whitespace = _shortcut_whitespace(column, identifier=identifier)
            else:
                whitespace = accumulator.finalise(identifier=identifier)
            whitespace_aggregates.append(
                {
                    "logical_table_id": table_id,
                    "source_column": rule.source_column,
                    "occurrence": rule.occurrence,
                    "qualified_name": rule.qualified_name,
                    **whitespace,
                    "raw_values_or_digests_included": False,
                }
            )

        declared_key_columns = list(
            table.get("declared_key_analysis", {}).get("columns", [])
        )
        if declared_key_columns:
            complete_rows = int(table["row_count"]) - key_blank_rows
            before_count = len(key_before)
            after_count = len(key_after)
            before_unique_complete = before_count == complete_rows
            after_unique_complete = after_count == complete_rows
            before_enforceable = key_blank_rows == 0 and before_count == int(table["row_count"])
            after_enforceable = key_blank_rows == 0 and after_count == int(table["row_count"])
            key_trim_analyses.append(
                {
                    "logical_table_id": table_id,
                    "declared_key_columns": declared_key_columns,
                    "business_primary_key_enforced": False,
                    "rows_with_any_blank_before_trim": key_blank_rows,
                    "rows_with_any_blank_after_trim": key_blank_rows,
                    "complete_rows": complete_rows,
                    "distinct_complete_combinations_before_trim": before_count,
                    "distinct_complete_combinations_after_trim": after_count,
                    "duplicate_complete_combinations_before_trim": complete_rows - before_count,
                    "duplicate_complete_combinations_after_trim": complete_rows - after_count,
                    "unique_when_complete_before_trim": before_unique_complete,
                    "unique_when_complete_after_trim": after_unique_complete,
                    "enforceable_not_null_unique_before_trim": before_enforceable,
                    "enforceable_not_null_unique_after_trim": after_enforceable,
                    "identifier_uniqueness_changed": (
                        before_unique_complete != after_unique_complete
                        or before_enforceable != after_enforceable
                    ),
                    "automatic_trim_applied": False,
                }
            )
        else:
            key_trim_analyses.append(
                {
                    "logical_table_id": table_id,
                    "declared_key_columns": [],
                    "status": "NOT_DECLARED_IN_SCHEMA",
                    "business_primary_key_enforced": False,
                    "identifier_uniqueness_changed": False,
                    "automatic_trim_applied": False,
                }
            )

        lineage_tables.append(
            {
                "logical_table_id": table_id,
                "physical_filename": source.physical_filename,
                "source_file_sha256": source.sha256,
                "source_sheet": str(table["selected_data_sheet"]),
                "header_excel_row": int(table["header_row"]),
                "first_data_excel_row": first_excel_row,
                "last_data_excel_row": last_excel_row,
                "source_row_count": int(table["row_count"]),
                "source_row_id_count": len(source_row_ids),
                "source_row_id_collision_count": int(table["row_count"]) - len(source_row_ids),
                "first_source_row_id": first_source_row_id,
                "last_source_row_id": last_source_row_id,
                "source_row_id_version": SOURCE_ROW_ID_VERSION,
                "business_primary_key_enforced": False,
            }
        )

    return {
        "parse_aggregates": parse_aggregates,
        "whitespace_aggregates": whitespace_aggregates,
        "lineage_tables": lineage_tables,
        "declared_key_trim_analyses": key_trim_analyses,
        "observed_total_rows": total_rows,
    }


__all__ = [
    "TypedStagingScanError",
    "build_rules_from_profiles",
    "scan_workbooks",
]
